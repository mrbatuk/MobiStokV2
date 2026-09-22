# -*- coding: utf-8 -*-
"""
Telefoncu uygulaması için 2 yıllık demo veri üretir.
Çıktı: demo_veriler.xlsx  (Ayarlar > Excel'den İçe Aktar ile yüklenebilir)

Tek sayfa "Cihazlar". İlk satır başlık; sütunlar başlık adına göre eşlenir.
Kolonlar:
  Kategori | Alış Tarihi | Model | IMEI | Alış Fiyatı | Satıcı | Genel Not |
  Alış Notu | Satış Tarihi | Satış Fiyatı | Alıcı | Satış Notu

Kategoriler: SIFIR, 2.EL, AKSESUAR (AKSESUAR'da IMEI boştur).
Kâr ve Bekleme hesaplanan alanlardır; içe aktarımda yok sayıldığı için dosyaya konmaz.
"""

import random
import sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

random.seed(42)  # tekrar üretilebilir demo

BASLANGIC = date(2024, 7, 1)
BITIS = date(2026, 7, 20)

# ---- Model havuzları: (model, min_alis, max_alis, min_kar, max_kar) ----
SIFIR_MODELLER = [
    ("iPhone 16 Pro Max 256GB", 62000, 68000, 3500, 7000),
    ("iPhone 16 Pro 128GB",     55000, 60000, 3000, 6000),
    ("iPhone 16 128GB",         45000, 50000, 2500, 5000),
    ("iPhone 15 128GB",         38000, 43000, 2500, 5000),
    ("iPhone 15 Plus 128GB",    42000, 46000, 2500, 5000),
    ("Samsung Galaxy S24 Ultra",48000, 54000, 3000, 6000),
    ("Samsung Galaxy S24",      32000, 37000, 2000, 4500),
    ("Samsung Galaxy A55",      16000, 19000, 1200, 2800),
    ("Samsung Galaxy A35",      12000, 14000, 1000, 2200),
    ("Samsung Galaxy A15",       7000,  8500,  700, 1600),
    ("Xiaomi Redmi Note 13 Pro",11000, 13000, 1000, 2200),
    ("Xiaomi Redmi Note 13",     8000,  9500,  800, 1800),
    ("Xiaomi 14",               28000, 32000, 1800, 4000),
    ("Tecno Spark 20",           4500,  5500,  500, 1200),
    ("Oppo Reno 11",            17000, 20000, 1200, 2800),
]

IKINCI_EL_MODELLER = [
    ("iPhone 13 128GB",        22000, 27000, 1500, 4000),
    ("iPhone 12 128GB",        16000, 20000, 1200, 3200),
    ("iPhone 11 64GB",         11000, 14000, 1000, 2800),
    ("iPhone XR 64GB",          8000, 10000,  800, 2200),
    ("iPhone SE 2020",          6000,  8000,  600, 1800),
    ("Samsung Galaxy S22",     15000, 18000, 1200, 3000),
    ("Samsung Galaxy S21",     11000, 14000, 1000, 2600),
    ("Samsung Galaxy A52",      6000,  8000,  600, 1700),
    ("Samsung Galaxy A32",      4500,  6000,  500, 1400),
    ("Xiaomi Mi 11",            8000, 10000,  800, 2000),
    ("Xiaomi Redmi Note 11",    5000,  6500,  500, 1500),
    ("Huawei P30 Lite",         3500,  5000,  400, 1300),
]

AKSESUAR_MODELLER = [
    ("Silikon Kılıf",            80,   250,   50,  200),
    ("Şeffaf Kılıf",             60,   180,   40,  150),
    ("Ekran Koruyucu Cam",      100,   300,   80,  250),
    ("20W Şarj Aleti",          250,   600,  120,  400),
    ("USB-C Kablo",             120,   350,   80,  300),
    ("Lightning Kablo",         150,   400,  100,  350),
    ("Bluetooth Kulaklık",      600,  1800,  300,  900),
    ("Kablolu Kulaklık",        150,   500,  100,  400),
    ("10000mAh Powerbank",      700,  1600,  300,  800),
    ("Araç Şarj Aleti",         200,   500,  120,  350),
]

SATICILAR = [
    "Vatan Toptan", "Teknosa Distribütör", "Mehmet Bey", "Ahmet Kaya",
    "Bireysel Müşteri", "İstanbul Toptancı", "Reyhan Elektronik",
    "Yılmaz Telekom", "Özkan Ticaret", "Bursa GSM",
]

ALICILAR = [
    "Ali Demir", "Ayşe Yıldız", "Fatih Şahin", "Zeynep Kaya", "Murat Aslan",
    "Elif Çelik", "Hasan Doğan", "Emine Arslan", "Burak Koç", "Seda Yılmaz",
    "Okan Polat", "Derya Öztürk", "Serkan Aydın", "Merve Kurt", "",  # bazıları isimsiz
]

GENEL_NOTLAR = ["", "", "", "", "Kutulu", "Faturalı", "Garantili", "Ekranda çizik var",
                "Aksesuarları tam", "Az kullanılmış", "Takas ile geldi"]
ALIS_NOTLARI = ["", "", "", "Peşin alındı", "Vadeli ödeme", "Toptan parti",
                "Servis çıkışlı", "Pazarlıkla alındı"]
SATIS_NOTLARI = ["", "", "", "", "Nakit satış", "Kredi kartı", "Taksitli",
                 "Kapıda ödeme", "Havale ile"]


def rastgele_gun(baslangic, bitis):
    return baslangic + timedelta(days=random.randint(0, (bitis - baslangic).days))


def gg_aa_yyyy(d):
    return d.strftime("%d.%m.%Y")


def imei_uret():
    return "".join(str(random.randint(0, 9)) for _ in range(15))


def enflasyon_carpani(alis_gun):
    # Zamanla fiyatların hafifçe artması (yıllık ~%15)
    ay_farki = (alis_gun.year - BASLANGIC.year) * 12 + (alis_gun.month - BASLANGIC.month)
    return 1 + (ay_farki / 12) * 0.15


def satirlar_uret(kategori, modeller, aylik_adet, imei_var, imza_seti):
    satirlar = []
    ay = date(BASLANGIC.year, BASLANGIC.month, 1)
    while ay <= BITIS:
        if ay.month == 12:
            sonraki = date(ay.year + 1, 1, 1)
        else:
            sonraki = date(ay.year, ay.month + 1, 1)
        ay_bitis = min(sonraki - timedelta(days=1), BITIS)

        adet = random.randint(*aylik_adet)
        for _ in range(adet):
            model, amin, amax, kmin, kmax = random.choice(modeller)
            alis_gun = rastgele_gun(ay, ay_bitis)
            carpan = enflasyon_carpani(alis_gun)
            birim = 50 if amax > 1000 else 5
            alis = round(random.randint(amin, amax) * carpan / birim) * birim

            if imei_var:
                imei = imei_uret()
                while imei in imza_seti:
                    imei = imei_uret()
                imza_seti.add(imei)
            else:
                imei = ""

            # Satış: çoğu satılmış; son 60 günde alınanların bir kısmı stokta
            bugune_kalan = (BITIS - alis_gun).days
            if bugune_kalan < 60:
                stokta_kal = random.random() < 0.55
            else:
                stokta_kal = random.random() < 0.08

            if stokta_kal:
                satis_gun = ""
                satis_fiyat = ""
                alici = ""
                satis_notu = ""
            else:
                gecikme = random.randint(1, 75)
                sg = alis_gun + timedelta(days=gecikme)
                if sg > BITIS:
                    sg = BITIS
                satis_gun = gg_aa_yyyy(sg)
                kar = round(random.randint(kmin, kmax) * carpan / birim) * birim
                satis_fiyat = int(alis + kar)
                alici = random.choice(ALICILAR)
                satis_notu = random.choice(SATIS_NOTLARI)

            satirlar.append([
                kategori,
                gg_aa_yyyy(alis_gun),
                model,
                imei,
                int(alis),
                random.choice(SATICILAR),
                random.choice(GENEL_NOTLAR),
                random.choice(ALIS_NOTLARI),
                satis_gun,
                satis_fiyat,
                alici,
                satis_notu,
            ])
        ay = sonraki
    return satirlar


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cihazlar"

    basliklar = ["Kategori", "Alış Tarihi", "Model", "IMEI", "Alış Fiyatı", "Satıcı",
                 "Genel Not", "Alış Notu", "Satış Tarihi", "Satış Fiyatı", "Alıcı", "Satış Notu"]
    ws.append(basliklar)
    baslik_font = Font(bold=True, color="FFFFFF")
    baslik_dolgu = PatternFill("solid", fgColor="2A78D6")
    for c in range(1, len(basliklar) + 1):
        h = ws.cell(row=1, column=c)
        h.font = baslik_font
        h.fill = baslik_dolgu
        h.alignment = Alignment(horizontal="center")

    imza_seti = set()
    tum = []
    tum += satirlar_uret("SIFIR", SIFIR_MODELLER, (14, 22), True, imza_seti)
    tum += satirlar_uret("2.EL", IKINCI_EL_MODELLER, (10, 16), True, imza_seti)
    tum += satirlar_uret("AKSESUAR", AKSESUAR_MODELLER, (8, 14), False, imza_seti)

    # Tarihe göre sırala (gg.aa.yyyy -> yyyyaagg)
    tum.sort(key=lambda r: r[1].split(".")[::-1])
    for r in tum:
        ws.append(r)

    genislik = [12, 13, 26, 18, 13, 20, 18, 18, 13, 13, 18, 18]
    for i, w in enumerate(genislik, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    cikti = sys.argv[1] if len(sys.argv) > 1 else "demo_veriler.xlsx"
    wb.save(cikti)

    def say(kat):
        alt = [r for r in tum if r[0] == kat]
        satilan = sum(1 for r in alt if r[8])
        return len(alt), satilan

    print(f"Olusturuldu: {cikti}  (tek sayfa: Cihazlar, toplam {len(tum)} kayit)")
    for kat in ("SIFIR", "2.EL", "AKSESUAR"):
        adet, satilan = say(kat)
        print(f"  {kat:9s}: {adet} kayit ({satilan} satilmis, {adet - satilan} stokta)")
    print(f"  Tarih araligi: {gg_aa_yyyy(BASLANGIC)} - {gg_aa_yyyy(BITIS)}")


if __name__ == "__main__":
    main()
